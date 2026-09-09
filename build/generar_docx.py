# -*- coding: utf-8 -*-
"""
Genera el .docx final de la monografía a partir de los archivos Markdown.

    python build/generar_docx.py

Los .md son la fuente de verdad; el .docx es un artefacto derivado y
descartable. Nunca editar el .docx a mano salvo para la revisión final.

Por qué python-docx y no Pandoc
-------------------------------
La Resolución N.º 1055/2024 exige tres cosas que Pandoc no puede expresar
desde Markdown y que igual habría que post-procesar con python-docx:

  1. Numeración romana en la parte pre textual y arábiga desde la Introducción
     (requiere saltos de sección con <w:pgNumType> propio).
  2. Una carátula con bloques en 24 pt conviviendo con bloques en 12 pt.
  3. Un campo TOC nativo de Word en el lugar del índice.

Hacer todo en un solo paso elimina la dependencia de Pandoc y deja una única
definición del formato, en build/formato.py.
"""

import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

from docx import Document                                    # noqa: E402
from docx.enum.section import WD_SECTION                     # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH                # noqa: E402
from docx.oxml import OxmlElement                            # noqa: E402
from docx.oxml.ns import qn                                  # noqa: E402
from docx.shared import Pt                                   # noqa: E402

from formato import (                                        # noqa: E402
    EXTENSION_MAXIMA_PAGINAS,
    EXTENSION_MINIMA_PAGINAS,
    FORMATO_NUM_CUERPO,
    FORMATO_NUM_PRELIMINARES,
    FUENTE,
    PALABRAS_POR_PAGINA,
    TAM_CARATULA,
    TAM_CUERPO,
    aplicar_fuente,
    campo,
    numerar_paginas,
)

RAIZ = Path(__file__).resolve().parent.parent
PLANTILLA = RAIZ / "referencia" / "plantilla_formato.docx"
SALIDA = RAIZ / "output" / "monografia-aranceles-consulares.docx"
COMBINADO = RAIZ / "output" / "_combinado.md"


# --------------------------------------------------------------------------
# Orden del documento
#
# Modificar SOLO acá para agregar, quitar o reordenar secciones.
# El tercer elemento indica si el bloque cuenta para el límite de extensión
# (20-50 páginas), que según la Resolución excluye la parte pre textual,
# la bibliografía y los anexos.
# --------------------------------------------------------------------------

CARATULA = ("preliminares/caratula.md", "caratula", False)

PRELIMINARES = [
    ("preliminares/dedicatoria.md", "dedicatoria", False),
    ("preliminares/agradecimiento.md", "agradecimiento", False),
    ("preliminares/indice.md", "indice", False),
]

CUERPO = [
    ("cuerpo/introduccion.md", "cuerpo", True),
    ("cuerpo/objetivos.md", "cuerpo", True),
    ("cuerpo/justificacion.md", "cuerpo", True),
    ("cuerpo/metodologia.md", "cuerpo", True),
    ("cuerpo/limitaciones.md", "cuerpo", True),
    ("cuerpo/capitulo-1.md", "cuerpo", True),
    ("cuerpo/capitulo-2.md", "cuerpo", True),
    ("cuerpo/capitulo-3.md", "cuerpo", True),
    ("cuerpo/capitulo-4.md", "cuerpo", True),
    ("cuerpo/capitulo-5.md", "cuerpo", True),
    ("cuerpo/conclusiones.md", "cuerpo", True),
    ("postexto/bibliografia.md", "bibliografia", False),
]


# --------------------------------------------------------------------------
# Lectura de Markdown
# --------------------------------------------------------------------------

RE_FRONTMATTER = re.compile(r"\A---\r?\n(.*?)\r?\n---[ \t]*\r?\n", re.S)
RE_COMENTARIO = re.compile(r"<!--.*?-->", re.S)
# Solo se reconoce énfasis con asteriscos: el guion bajo aparece dentro de URLs
# de la bibliografía y tratarlo como cursiva las rompería.
RE_ENFASIS = re.compile(r"(\*\*.+?\*\*|\*[^*\n]+?\*)")


def leer_frontmatter(texto):
    """Analiza el bloque YAML de la carátula. Soporta claves simples y listas."""
    coincidencia = RE_FRONTMATTER.search(texto)
    if not coincidencia:
        return {}
    datos, clave_lista = {}, None
    for linea in coincidencia.group(1).splitlines():
        desnuda = linea.strip()
        if not desnuda or desnuda.startswith("#"):
            continue
        if desnuda.startswith("- "):
            if clave_lista:
                datos[clave_lista].append(desnuda[2:].strip().strip('"\''))
            continue
        if ":" in desnuda:
            clave, _, valor = desnuda.partition(":")
            clave, valor = clave.strip(), valor.strip().strip('"\'')
            if valor:
                datos[clave] = valor
                clave_lista = None
            else:
                datos[clave] = []
                clave_lista = clave
    return datos


def analizar(texto):
    """Convierte Markdown en una lista de bloques (tipo, texto).

    Deliberadamente mínimo: reconoce lo que este trabajo usa —títulos de tres
    niveles, párrafos, viñetas y citas en bloque— y nada más. Los comentarios
    HTML (notas de trabajo) se descartan y no llegan al .docx.
    """
    texto = RE_FRONTMATTER.sub("", texto)
    texto = RE_COMENTARIO.sub("", texto)

    bloques, acumulado = [], []

    def cerrar_parrafo():
        if acumulado:
            bloques.append(("parrafo", " ".join(acumulado).strip()))
            acumulado.clear()

    for linea in texto.splitlines():
        desnuda = linea.strip()
        if not desnuda:
            cerrar_parrafo()
        elif desnuda.startswith("### "):
            cerrar_parrafo()
            bloques.append(("h3", desnuda[4:].strip()))
        elif desnuda.startswith("## "):
            cerrar_parrafo()
            bloques.append(("h2", desnuda[3:].strip()))
        elif desnuda.startswith("# "):
            cerrar_parrafo()
            bloques.append(("h1", desnuda[2:].strip()))
        elif desnuda.startswith("> "):
            cerrar_parrafo()
            bloques.append(("cita", desnuda[2:].strip()))
        elif desnuda[:2] in ("- ", "* "):
            cerrar_parrafo()
            bloques.append(("vineta", desnuda[2:].strip()))
        else:
            acumulado.append(desnuda)
    cerrar_parrafo()
    return bloques


def contar_palabras(bloques):
    return sum(len(texto.split()) for tipo, texto in bloques if tipo != "h1")


# --------------------------------------------------------------------------
# Escritura en el documento
# --------------------------------------------------------------------------

def escribir_texto(parrafo, texto, tam=TAM_CUERPO):
    """Vuelca texto con énfasis en negrita/cursiva resueltos como runs."""
    for fragmento in RE_ENFASIS.split(texto):
        if not fragmento:
            continue
        if fragmento.startswith("**") and fragmento.endswith("**"):
            run = parrafo.add_run(fragmento[2:-2])
            run.bold = True
        elif fragmento.startswith("*") and fragmento.endswith("*"):
            run = parrafo.add_run(fragmento[1:-1])
            run.italic = True
        else:
            run = parrafo.add_run(fragmento)
        aplicar_fuente(run, FUENTE, tam)
    return parrafo


def agregar_titulo(documento, nivel, texto, salto_pagina):
    if nivel == 1:
        # "centrado, mayúsculas, negrita, sin punto final".
        texto = texto.rstrip(".").upper()
    parrafo = documento.add_paragraph(style=f"Heading {nivel}")
    parrafo.paragraph_format.page_break_before = salto_pagina
    return escribir_texto(parrafo, texto)


def agregar_bloques(documento, bloques, estilo_parrafo="Normal",
                    estilo_vineta="List Bullet", primer_salto=False):
    """Escribe una lista de bloques. Cada título de nivel 1 abre página nueva."""
    es_primero = True
    for tipo, texto in bloques:
        if tipo == "h1":
            agregar_titulo(documento, 1, texto,
                           salto_pagina=primer_salto if es_primero else True)
        elif tipo in ("h2", "h3"):
            agregar_titulo(documento, int(tipo[1]), texto, salto_pagina=False)
        elif tipo == "vineta":
            escribir_texto(documento.add_paragraph(style=estilo_vineta), texto)
        elif tipo == "cita":
            escribir_texto(documento.add_paragraph(style="Cita en bloque"), texto)
        else:
            escribir_texto(documento.add_paragraph(style=estilo_parrafo), texto)
        es_primero = False


def espaciador(documento, cantidad=1):
    for _ in range(cantidad):
        parrafo = documento.add_paragraph()
        parrafo.paragraph_format.space_after = Pt(0)


# --------------------------------------------------------------------------
# Piezas especiales
# --------------------------------------------------------------------------

def agregar_caratula(documento, datos):
    """Portada según la Resolución N.º 1055/2024, parte pre textual, punto 1."""
    def linea(texto, tam):
        estilo = "Caratula 24" if tam == TAM_CARATULA else "Caratula 12"
        parrafo = documento.add_paragraph(style=estilo)
        return escribir_texto(parrafo, texto, tam=tam)

    espaciador(documento, 1)
    for texto in datos.get("institucion_24pt", []):
        linea(texto, TAM_CARATULA)

    espaciador(documento, 2)
    linea(datos.get("curso", ""), TAM_CUERPO)
    linea(datos.get("fundamento", ""), TAM_CUERPO)

    espaciador(documento, 2)
    linea(datos.get("titulo_24pt", ""), TAM_CARATULA)

    espaciador(documento, 2)
    linea(datos.get("autor", ""), TAM_CUERPO)

    espaciador(documento, 2)
    linea(datos.get("ciudad", ""), TAM_CUERPO)
    linea(datos.get("mes_anio_24pt", ""), TAM_CARATULA)


def agregar_indice(documento, bloques):
    """Encabezado ÍNDICE + campo TOC nativo, que Word actualiza solo."""
    for tipo, texto in bloques:
        if tipo == "h1":
            agregar_titulo(documento, 1, texto, salto_pagina=True)

    parrafo = documento.add_paragraph()
    parrafo.alignment = WD_ALIGN_PARAGRAPH.LEFT
    campo(
        parrafo,
        ' TOC \\o "1-3" \\h \\z \\u ',
        marcador="[Índice automático: colocar el cursor acá y presionar F9 "
                 "para generarlo.]",
    )


def numero_de_pagina(seccion, mostrar=True):
    """Numeración centrada en la parte superior, vía encabezado."""
    encabezado = seccion.header
    encabezado.is_linked_to_previous = False

    for parrafo in list(encabezado.paragraphs)[1:]:
        parrafo._element.getparent().remove(parrafo._element)

    parrafo = encabezado.paragraphs[0]
    for run in list(parrafo.runs):
        run._element.getparent().remove(run._element)
    parrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    parrafo.paragraph_format.space_after = Pt(0)

    if mostrar:
        campo(parrafo, " PAGE ", marcador="1")


def actualizar_campos_al_abrir(documento):
    """<w:updateFields/>: Word ofrece regenerar el índice al abrir el archivo."""
    ajustes = documento.settings.element
    if ajustes.find(qn("w:updateFields")) is not None:
        return
    elemento = OxmlElement("w:updateFields")
    elemento.set(qn("w:val"), "true")
    # Debe preceder a w:compat / w:rsids para no violar el orden del esquema.
    for etiqueta in ("w:compat", "w:rsids", "w:mathPr", "w:themeFontLang"):
        anclaje = ajustes.find(qn(etiqueta))
        if anclaje is not None:
            anclaje.addprevious(elemento)
            return
    ajustes.append(elemento)


# --------------------------------------------------------------------------
# Construcción
# --------------------------------------------------------------------------

def cargar(ruta_relativa):
    ruta = RAIZ / ruta_relativa
    if not ruta.exists():
        raise SystemExit(f"Falta el archivo declarado en el manifiesto: {ruta_relativa}")
    return ruta.read_text(encoding="utf-8")


def generar():
    if not PLANTILLA.exists():
        from crear_plantilla import construir_plantilla
        construir_plantilla()
        print(f"Plantilla ausente: se generó {PLANTILLA.relative_to(RAIZ)}")

    documento = Document(str(PLANTILLA))
    for parrafo in list(documento.paragraphs):
        parrafo._element.getparent().remove(parrafo._element)

    crudo = {}
    analizados = {}
    for ruta, _, _ in [CARATULA] + PRELIMINARES + CUERPO:
        crudo[ruta] = cargar(ruta)
        analizados[ruta] = analizar(crudo[ruta])

    # --- Sección 1: carátula. Sin número de página visible. ---
    seccion_caratula = documento.sections[0]
    numero_de_pagina(seccion_caratula, mostrar=False)
    numerar_paginas(seccion_caratula, FORMATO_NUM_PRELIMINARES, inicio=1)
    agregar_caratula(documento, leer_frontmatter(crudo[CARATULA[0]]))

    # --- Sección 2: resto de la parte pre textual, en números romanos. ---
    seccion_preliminares = documento.add_section(WD_SECTION.NEW_PAGE)
    numero_de_pagina(seccion_preliminares, mostrar=True)
    # La carátula ocupa la página i, que no se imprime: esta sección arranca en ii.
    numerar_paginas(seccion_preliminares, FORMATO_NUM_PRELIMINARES, inicio=2)

    for ruta, modo, _ in PRELIMINARES:
        bloques = analizados[ruta]
        if modo == "indice":
            agregar_indice(documento, bloques)
        elif modo == "dedicatoria":
            agregar_titulo(documento, 1, bloques[0][1], salto_pagina=False)
            espaciador(documento, 3)
            agregar_bloques(documento, bloques[1:], estilo_parrafo="Dedicatoria")
        else:
            agregar_titulo(documento, 1, bloques[0][1], salto_pagina=True)
            agregar_bloques(documento, bloques[1:], estilo_parrafo="Agradecimiento")

    # --- Sección 3: cuerpo y post-texto, en números arábigos desde 1. ---
    seccion_cuerpo = documento.add_section(WD_SECTION.NEW_PAGE)
    numero_de_pagina(seccion_cuerpo, mostrar=True)
    numerar_paginas(seccion_cuerpo, FORMATO_NUM_CUERPO, inicio=1)

    primero = True
    for ruta, modo, _ in CUERPO:
        estilo = "Bibliografia APA" if modo == "bibliografia" else "Normal"
        vineta = "Bibliografia APA" if modo == "bibliografia" else "List Bullet"
        agregar_bloques(documento, analizados[ruta], estilo_parrafo=estilo,
                        estilo_vineta=vineta, primer_salto=not primero)
        primero = False

    actualizar_campos_al_abrir(documento)

    SALIDA.parent.mkdir(parents=True, exist_ok=True)
    documento.save(str(SALIDA))

    COMBINADO.write_text(
        "\n\n".join(crudo[ruta] for ruta, _, _ in PRELIMINARES + CUERPO),
        encoding="utf-8",
    )
    return analizados


def informar(analizados):
    print(f"\nDocumento generado: {SALIDA.relative_to(RAIZ)}\n")

    palabras_computables = 0
    print("Extensión por archivo (palabras):")
    for ruta, _, computa in CUERPO:
        palabras = contar_palabras(analizados[ruta])
        if computa:
            palabras_computables += palabras
        marca = " " if computa else "*"
        estado = "PENDIENTE" if palabras == 0 else ""
        print(f"  {marca} {ruta:<32} {palabras:>6}  {estado}")
    print("  (* no computa para el límite de extensión)")

    paginas = palabras_computables / PALABRAS_POR_PAGINA
    print(f"\nCuerpo computable: {palabras_computables} palabras "
          f"(~{paginas:.1f} páginas estimadas)")
    if paginas < EXTENSION_MINIMA_PAGINAS:
        faltan = (EXTENSION_MINIMA_PAGINAS - paginas) * PALABRAS_POR_PAGINA
        print(f"  -> Por debajo del mínimo de {EXTENSION_MINIMA_PAGINAS} páginas: "
              f"faltan ~{faltan:.0f} palabras.")
    elif paginas > EXTENSION_MAXIMA_PAGINAS:
        print(f"  -> Por encima del máximo de {EXTENSION_MAXIMA_PAGINAS} páginas.")
    else:
        print("  -> Dentro del rango exigido.")

    capitulos = [(ruta, contar_palabras(analizados[ruta]))
                 for ruta, _, _ in CUERPO if "capitulo-" in ruta]
    escritos = [palabras for _, palabras in capitulos if palabras]
    if len(escritos) > 1:
        print(f"\nEquilibrio entre capítulos redactados: "
              f"menor {min(escritos)} / mayor {max(escritos)} palabras")

    print("\nAl abrir en Word: aceptar la actualización de campos para que se "
          "construya el índice.\nSi no aparece el aviso: Ctrl+E y luego F9.")


if __name__ == "__main__":
    informar(generar())
