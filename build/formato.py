# -*- coding: utf-8 -*-
"""
Constantes y utilidades de formato comunes.

TODAS las reglas de la Resolución N.º 1055/2024 (Anexo) que sean numéricas o
tipográficas viven acá, en un solo lugar, para que cualquier ajuste posterior
sea un cambio de una línea y no una cacería por el código.

Referencia: Resolución N.º 1055/2024 de la Academia Diplomática y Consular
"Carlos Antonio López", Anexo, secciones "Características formales generales",
"Parte pre textual", "Cuerpo del trabajo" y "Post-texto".
"""

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

# --------------------------------------------------------------------------
# 3.1 Características formales generales
# --------------------------------------------------------------------------

FUENTE = "Times New Roman"

TAM_CUERPO = 12          # cuerpo del trabajo
TAM_NOTA_PIE = 10        # notas al pie
TAM_CARATULA = 24        # bloques de la carátula indicados en la Resolución

INTERLINEADO = 1.5

# "Espacio libre entre párrafos: 1,5 (recomendado)".
# Se interpreta como un espacio equivalente a 1,5 líneas del cuerpo
# (1,5 x 12 pt = 18 pt) después de cada párrafo.
ESPACIO_ENTRE_PARRAFOS = Pt(18)

# Hoja A4 e impresión a una sola carilla.
ANCHO_HOJA = Cm(21.0)
ALTO_HOJA = Cm(29.7)

MARGEN_SUPERIOR = Cm(2.5)
MARGEN_INFERIOR = Cm(2.5)
MARGEN_DERECHO = Cm(2.5)
MARGEN_IZQUIERDO = Cm(3.5)

# La Resolución no se pronuncia sobre la alineación del texto corrido.
# Se adopta la justificación, de uso estándar en los trabajos monográficos
# de la Academia. Cambiar a WD_ALIGN_PARAGRAPH.LEFT si el Tribunal lo objeta.
ALINEACION_CUERPO = WD_ALIGN_PARAGRAPH.JUSTIFY

# Sangría francesa de las referencias bibliográficas (APA 7.ª ed., 2019).
SANGRIA_BIBLIOGRAFIA = Cm(1.27)

# Sangría de las citas textuales en bloque (más de 40 palabras, APA 7).
SANGRIA_CITA = Cm(1.27)

# 3.1 "Numeración de páginas: parte pre textual en números romanos".
# La Resolución no distingue entre mayúsculas y minúsculas; se usa la forma
# minúscula (i, ii, iii), la más habitual. Cambiar a "upperRoman" si se
# prefiere I, II, III.
FORMATO_NUM_PRELIMINARES = "lowerRoman"
FORMATO_NUM_CUERPO = "decimal"

IDIOMA = "es-PY"

# Estimación usada solo para el informe de extensión que imprime el build.
# No sustituye al recuento real de páginas que hace Word.
PALABRAS_POR_PAGINA = 380
EXTENSION_MINIMA_PAGINAS = 20
EXTENSION_MAXIMA_PAGINAS = 50


# --------------------------------------------------------------------------
# Utilidades XML de bajo nivel
# --------------------------------------------------------------------------

# Orden de los hijos de <w:sectPr> según el esquema WordprocessingML.
# Insertar fuera de este orden hace que Word marque el documento como dañado.
_ORDEN_SECTPR = (
    "w:headerReference", "w:footerReference", "w:footnotePr", "w:endnotePr",
    "w:type", "w:pgSz", "w:pgMar", "w:paperSrc", "w:pgBorders", "w:lnNumType",
    "w:pgNumType", "w:cols", "w:formProt", "w:vAlign", "w:noEndnote",
    "w:titlePg", "w:textDirection", "w:bidi", "w:rtlGutter", "w:docGrid",
    "w:printerSettings", "w:sectPrChange",
)


def insertar_en_orden(padre, elemento, etiqueta, orden):
    """Inserta `elemento` entre los hijos de `padre` respetando `orden`."""
    posteriores = {qn(t) for t in orden[orden.index(etiqueta) + 1:]}
    for hijo in padre:
        if hijo.tag in posteriores:
            hijo.addprevious(elemento)
            return
    padre.append(elemento)


def numerar_paginas(seccion, formato, inicio=None):
    """Fija el formato (romano / arábigo) y el número inicial de una sección.

    Word no permite expresar esto desde Markdown: es la razón principal por la
    que el build no se apoya en Pandoc.
    """
    sect_pr = seccion._sectPr
    pg_num = sect_pr.find(qn("w:pgNumType"))
    if pg_num is None:
        pg_num = OxmlElement("w:pgNumType")
        insertar_en_orden(sect_pr, pg_num, "w:pgNumType", _ORDEN_SECTPR)
    pg_num.set(qn("w:fmt"), formato)
    if inicio is not None:
        pg_num.set(qn("w:start"), str(inicio))


def aplicar_fuente(fuente_o_run, nombre=FUENTE, tam=None):
    """Aplica la familia tipográfica en las cuatro variantes que usa Word.

    Sin fijar w:eastAsia y w:cs, Word puede sustituir la fuente en caracteres
    acentuados o en símbolos, rompiendo la uniformidad exigida.
    """
    objeto = getattr(fuente_o_run, "font", fuente_o_run)
    objeto.name = nombre
    if tam is not None:
        objeto.size = Pt(tam)

    elemento = getattr(fuente_o_run, "element", None) or fuente_o_run._element
    r_pr = elemento.get_or_add_rPr()
    r_fonts = r_pr.get_or_add_rFonts()
    for atributo in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        r_fonts.set(qn(atributo), nombre)


def campo(parrafo, instruccion, marcador=""):
    """Inserta un campo nativo de Word (PAGE, TOC, ...) en un párrafo.

    `marcador` es el texto que se ve mientras el campo no está actualizado.
    """
    run_inicio = parrafo.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    run_inicio._r.append(fld_begin)

    run_instr = parrafo.add_run()
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruccion
    run_instr._r.append(instr)

    run_sep = parrafo.add_run()
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    run_sep._r.append(fld_sep)

    run_marcador = parrafo.add_run(marcador)

    run_fin = parrafo.add_run()
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run_fin._r.append(fld_end)

    for run in (run_inicio, run_instr, run_sep, run_marcador, run_fin):
        aplicar_fuente(run, FUENTE, TAM_CUERPO)
    return parrafo
