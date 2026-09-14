# -*- coding: utf-8 -*-
"""
Verificación estructural del .docx generado, contra la Resolución N.º 1055/2024.

    python build/verificar.py [ruta_al_docx]

Lee el .docx ya construido y comprueba, una por una, las reglas del Anexo:
secciones y numeración de páginas, medidas de la hoja, estilos, campos, y el
orden y la caja de los títulos. No abre Word: trabaja sobre el XML del archivo.
Para verificar lo que Word efectivamente compone —índice actualizado, recuento
real de páginas— está `build/verificar_en_word.ps1`.

Devuelve 0 si todo pasa, 1 si algo falla.
"""

import sys
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Cm

RAIZ = Path(__file__).resolve().parent.parent
POR_DEFECTO = RAIZ / "output" / "monografia-aranceles-consulares.docx"

# Cantidades esperadas del documento en su estado actual. Actualizar cuando el
# contenido cambie a propósito: que estos números fallen es la señal de que algo
# se agregó o se perdió sin querer.
ENTRADAS_BIBLIOGRAFIA = 16   # 9 fuentes primarias + 7 secundarias
VINETAS_OBJETIVOS = 6        # 1 objetivo general + 5 específicos
CANTIDAD_CAPITULOS = 5

resultados = []


def chk(condicion, mensaje):
    resultados.append(bool(condicion))
    print(("  OK   " if condicion else "  FALLA") + "  " + mensaje)


def verificar(ruta):
    doc = Document(str(ruta))

    print("=== SECCIONES Y NUMERACIÓN ===")
    # (formato, página inicial, ¿lleva número impreso?)
    esperado = [
        ("lowerRoman", "1", False),   # carátula: es la página i, pero no se imprime
        ("lowerRoman", "2", True),    # dedicatoria, agradecimientos, índice
        ("decimal", "1", True),       # de la Introducción en adelante
    ]
    chk(len(doc.sections) == 3, "3 secciones (carátula / preliminares / cuerpo)")
    for i, seccion in enumerate(doc.sections):
        pg_num = seccion._sectPr.find(qn("w:pgNumType"))
        fmt = pg_num.get(qn("w:fmt")) if pg_num is not None else None
        inicio = pg_num.get(qn("w:start")) if pg_num is not None else None
        tiene_campo = "PAGE" in seccion.header.paragraphs[0]._p.xml
        e_fmt, e_inicio, e_campo = esperado[i]
        chk(fmt == e_fmt and inicio == e_inicio,
            f"sección {i}: fmt={fmt} start={inicio} (esperado {e_fmt}/{e_inicio})")
        chk(tiene_campo == e_campo,
            f"sección {i}: número de página en el encabezado = {tiene_campo}")
        if e_campo:
            chk(str(seccion.header.paragraphs[0].alignment) == "CENTER (1)",
                f"sección {i}: encabezado centrado")

    print("\n=== HOJA Y MÁRGENES ===")
    s = doc.sections[0]
    for etiqueta, real, nominal in (
        ("ancho A4", s.page_width, 21.0),
        ("alto A4", s.page_height, 29.7),
        ("margen superior", s.top_margin, 2.5),
        ("margen inferior", s.bottom_margin, 2.5),
        ("margen derecho", s.right_margin, 2.5),
        ("margen izquierdo", s.left_margin, 3.5),
    ):
        chk(abs(real - Cm(nominal)) < 1000, f"{etiqueta} = {real.cm:.2f} cm")

    print("\n=== ESTILOS ===")
    normal = doc.styles["Normal"]
    chk(normal.font.name == "Times New Roman", f"Normal: {normal.font.name}")
    chk(normal.font.size.pt == 12, f"Normal: {normal.font.size.pt} pt")
    chk(abs(normal.paragraph_format.line_spacing - 1.5) < 0.01,
        f"Normal: interlineado {normal.paragraph_format.line_spacing}")
    chk(normal.paragraph_format.space_after.pt == 18,
        f"Normal: {normal.paragraph_format.space_after.pt} pt entre párrafos")

    for nombre, tam in (("Footnote Text", 10), ("Endnote Text", 10),
                        ("Caratula 24", 24), ("Caratula 12", 12)):
        estilo = doc.styles[nombre]
        chk(estilo.font.name == "Times New Roman" and estilo.font.size.pt == tam,
            f"{nombre}: {estilo.font.name} {estilo.font.size.pt} pt")

    for nombre in ("Heading 1", "Heading 2", "Heading 3"):
        estilo = doc.styles[nombre]
        chk(estilo.font.bold and estilo.font.name == "Times New Roman"
            and estilo.font.size.pt == 12,
            f"{nombre}: negrita, {estilo.font.name} {estilo.font.size.pt} pt")

    chk(str(doc.styles["Heading 1"].paragraph_format.alignment) == "CENTER (1)",
        "Heading 1 centrado")
    biblio = doc.styles["Bibliografia APA"]
    chk(biblio.paragraph_format.first_line_indent < 0,
        f"Bibliografía: sangría francesa de {biblio.paragraph_format.first_line_indent.cm:.2f} cm")
    chk(str(doc.styles["Dedicatoria"].paragraph_format.alignment) == "RIGHT (2)",
        "Dedicatoria al margen derecho")
    chk(str(doc.styles["Agradecimiento"].paragraph_format.alignment) == "CENTER (1)",
        "Agradecimiento centrado")
    for nombre in ("TOC 1", "TOC 2", "TOC 3"):
        chk(doc.styles[nombre].font.bold, f"{nombre} en negrita")

    print("\n=== CAMPOS ===")
    chk('TOC \\o "1-3"' in doc.element.xml, "campo TOC presente")
    chk(doc.settings.element.find(qn("w:updateFields")) is not None,
        "updateFields activado (Word ofrece armar el índice al abrir)")

    print("\n=== TÍTULOS DE NIVEL 1 ===")
    titulos = [p.text for p in doc.paragraphs if p.style.name == "Heading 1"]
    for texto in titulos:
        print("   -", texto)
    chk(all(t == t.upper() for t in titulos), "todos en mayúsculas")
    chk(all(not t.endswith(".") for t in titulos), "ninguno termina en punto")
    chk(titulos[:6] == ["DEDICATORIA", "AGRADECIMIENTOS", "INTRODUCCIÓN",
                        "JUSTIFICACIÓN", "METODOLOGÍA", "LIMITACIONES"],
        "orden de los primeros seis")
    chk(titulos[-1] == "BIBLIOGRAFÍA", f"el último es {titulos[-1]}")
    chk(sum(1 for t in titulos if t.startswith("CAPÍTULO")) == CANTIDAD_CAPITULOS,
        f"{CANTIDAD_CAPITULOS} capítulos")

    # El switch \o del campo TOC selecciona POR ESTILO, no por nivel de esquema:
    # por eso el encabezado ÍNDICE lleva un estilo propio y no "Heading 1".
    indice = [p for p in doc.paragraphs if p.text == "ÍNDICE"]
    chk(len(indice) == 1 and indice[0].style.name == "Titulo preliminar",
        "el encabezado ÍNDICE queda fuera del propio índice")
    preliminar = doc.styles["Titulo preliminar"]
    chk(preliminar.font.bold
        and str(preliminar.paragraph_format.alignment) == "CENTER (1)",
        "'Titulo preliminar' se ve igual que un título de nivel 1")

    print("\n=== CONTENIDO ===")
    textos = [p.text for p in doc.paragraphs]
    chk(not any("<!--" in t for t in textos), "sin comentarios HTML en el .docx")
    chk(not any("Pendiente de redacción" in t for t in textos),
        "las notas de trabajo no llegan al documento")
    chk(any("Ross Chamorro" in t for t in textos), "dedicatoria presente")
    chk(any(t.startswith("Propuesta de automatización") for t in textos),
        "título del trabajo en la carátula")
    chk(any(t == "Sven Knutson Sachelaridi" for t in textos), "autor en la carátula")
    chk(not any("*" in t for t in textos), "sin asteriscos de Markdown sin procesar")

    cursivas = [r.text for p in doc.paragraphs for r in p.runs if r.italic]
    chk("El fin del trámite eterno: ciudadanos, burocracia y gobierno digital" in cursivas,
        "las cursivas de Markdown se convirtieron")

    entradas = [p for p in doc.paragraphs if p.style.name == "Bibliografia APA"]
    chk(len(entradas) == ENTRADAS_BIBLIOGRAFIA,
        f"{len(entradas)} entradas de bibliografía (esperado {ENTRADAS_BIBLIOGRAFIA})")
    vinetas = [p for p in doc.paragraphs if p.style.name == "List Bullet"]
    chk(len(vinetas) == VINETAS_OBJETIVOS,
        f"{len(vinetas)} viñetas de objetivos (esperado {VINETAS_OBJETIVOS})")

    otras_fuentes = {r.font.name for p in doc.paragraphs for r in p.runs
                     if r.font.name not in (None, "Times New Roman")}
    chk(not otras_fuentes, f"todo en Times New Roman (otras: {otras_fuentes or 'ninguna'})")


if __name__ == "__main__":
    ruta = Path(sys.argv[1]) if len(sys.argv) > 1 else POR_DEFECTO
    if not ruta.exists():
        raise SystemExit(f"No existe {ruta}. Correr antes: python build/generar_docx.py")

    verificar(ruta)
    total, buenas = len(resultados), sum(resultados)
    print(f"\n=== RESULTADO: {buenas}/{total} comprobaciones OK ===")
    sys.exit(0 if buenas == total else 1)
